"""Create the team-authored technical DOCX; never generates an official form.

Offline authoring dependency: python-docx (separate from HomeTrace runtime).
Run with --test-count only after the corresponding real test log exists.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"


def text(document, value, style=None):
    return document.add_paragraph(value, style)


def table(document, headers, rows, widths):
    result = document.add_table(rows=1, cols=len(headers))
    result.autofit = False
    for index, width in enumerate(widths):
        result.columns[index].width = Cm(width)
    for index, value in enumerate(headers):
        result.rows[0].cells[index].text = value
    for row in rows:
        for cell, value in zip(result.add_row().cells, row):
            cell.text = str(value)
    for row_index, row in enumerate(result.rows):
        for column, cell in enumerate(row.cells):
            cell.width = Cm(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "17324D" if row_index == 0 else ("F1F4F6" if row_index % 2 == 0 else "FFFFFF"))
            props.append(shading)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}")
                for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    element.set(qn(f"w:{key}"), value)
                borders.append(element)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                element = OxmlElement(f"w:{edge}")
                element.set(qn("w:w"), "100")
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            props.append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = Pt(14)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_index == 0 else "000000")
                    run.bold = row_index == 0
    header = OxmlElement("w:tblHeader")
    result.rows[0]._tr.get_or_add_trPr().append(header)
    spacer = text(document, "")
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    return result


def heading(document, value):
    document.add_heading(value, level=1)


def picture(document, name, caption, width=15.8):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run().add_picture(str(ROOT / "docs/evidence/issue-3/browser" / name), width=Cm(width))
    caption_p = text(document, caption, "Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build(output: Path, test_count: int | None, test_log: str | None):
    if test_count is not None and (not test_log or not (ROOT / test_log).is_file()):
        raise ValueError("A real test log must exist before a current test count is printed")
    document = Document()
    for style in document.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        if style.type == 1:
            snap = OxmlElement("w:snapToGrid")
            snap.set(qn("w:val"), "0")
            style.element.get_or_add_pPr().append(snap)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    section.footer_distance = Cm(0.9)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
        style = document.styles[name]
        style.font.name = FONT
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = Pt(16)
    normal.paragraph_format.space_after = Pt(7)
    document.styles["Title"].font.size = Pt(27)
    document.styles["Title"].paragraph_format.line_spacing = 1.15
    document.styles["Title"].paragraph_format.space_after = Pt(18)
    document.styles["Heading 1"].font.size = Pt(17)
    document.styles["Heading 1"].paragraph_format.line_spacing = Pt(24)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(9)
    document.styles["Heading 2"].font.size = Pt(12)
    document.styles["Caption"].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("HomeTrace  ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "HomeTrace Team"
    document.core_properties.title = "HomeTrace 作品设计与技术说明书"
    document.core_properties.subject = "team_authored 技术说明书"
    document.core_properties.comments = "Public version excludes private photographs and personal identity fields."

    text(document, "HomeTrace\n作品设计与技术说明书", "Title")
    text(document, "家庭物品的时空记忆与主动寻找", "Subtitle")
    text(document, "2026 年华为嵌入式软件大赛初赛作品材料\n文档类型  team_authored 团队自编技术说明书\n团队名称  待团队填写\n队长与联系方式  待团队在本地填写")
    heading(document, "作品摘要")
    text(document, "HomeTrace 为已登记物品保存历史位置，并在新的搜索任务中读取当前图像、调整搜索计划和验证结果。作品把记忆、主动搜索与可追溯证据组成可运行的软件闭环。当前 Demo 在普通 Windows 计算机运行，使用 SQLite 和成熟 OpenCV 特征匹配，不依赖大型模型。")
    text(document, "稳定公开演示已覆盖找到目标、未确认负例和旧图重放保护。团队家庭实拍照片另行实测，当前未得到满足原验证条件的 FOUND，系统如实保留 SUSPECTED 或 UNKNOWN。所有导航为模拟，没有真实机器人运行数据。")
    heading(document, "验证结论")
    table(document, ["范围", "结果与含义"], [
        ["公开照片闭环", "正例 FOUND/B，负例 SUSPECTED，重放 UNKNOWN"],
        ["家庭图片适配", "75 次检测评估中 0 FOUND，详见第 5 页的样本口径与原因"],
        ["策略效率", "三个公开场景步数均打平，尚未证明效率提高"],
        ["代码回归", f"P3 本轮 {test_count} 项通过，日志见第 6 页" if test_count is not None else "P2 稳定基线 148 项通过，P3 本轮见运行测试报告"],
    ], [3.2, 13.4])

    document.add_page_break()
    heading(document, "1 用户问题与创新思路")
    text(document, "家庭成员常记得物品曾经放在哪里，却不能确认现在的位置。HomeTrace 把历史线索作为搜索起点，再用当前观察修正判断。用户只需选择已登记物品，系统解释为什么检查某个地点、看到了哪些证据，以及为什么确认或继续寻找。")
    text(document, "作品的创新点是把历史位置与当前证据分离，在搜索过程中持续更新判断，并保留完整决策链。历史记忆决定优先检查的候选点，不能绕过当前图像验证。一次匹配失败也不等于物品不存在，系统保留不确定性并继续探索候选点。")
    heading(document, "2 软件结构与实现")
    table(document, ["模块", "职责"], [
        ["Search Core", "查询历史、规划候选点、消费观察、重规划、输出结束原因"],
        ["SQLite 记忆", "跨进程保存登记物品、参考图摘要、历史与已见证据"],
        ["视觉适配", "读取实际像素，执行 ORB、Hamming 匹配与 RANSAC 几何验证"],
        ["导航接口", "返回模拟到达事件，明确 simulated_navigation 来源"],
        ["本地页面", "展示物品、历史、当前图、检测框、逐步事件与报告下载"],
    ], [3.2, 13.4])
    text(document, "Search Core 与视觉、界面和导航接口解耦。现阶段使用 Python 标准库 HTTP 服务与静态 HTML/JS，视觉为可选 OpenCV 依赖。未来可以替换观察来源或导航实现，但必须保留来源、时序与验证契约。当前没有 ROS、机器人 SDK 或设备接入。")
    heading(document, "3 可信结果的约束")
    text(document, "每次当前观察记录图片来源、SHA256、时间、匹配结果与失败原因。旧图即使更换名称或时间也不能成为新的确认。FOUND 需要当前图像和几何支持通过条件；SUSPECTED 表示有候选但尚未确认；UNKNOWN 表示本轮没有足以确认的证据。两种未确认状态都不证明目标不存在。")
    text(document, "FOUND 只验证照片中的登记外观。平面纹理和透视条件会影响识别，同款包装的多个实体不能可靠区分。物品真值标签、场景答案和用户指定位置均不能直接输出为视觉检测结果。")

    document.add_page_break()
    heading(document, "4 公开正例展示")
    text(document, "登记进程将目标与历史 A 写入 SQLite。新的搜索进程优先检查 A，当前照片未确认后继续到 B。B 的独立场景照片通过匹配与几何验证，最终输出 FOUND 和确认地点 B。")
    picture(document, "moved-matched.png", "图 1  真实浏览器截图  公开实拍照片  模拟导航")
    text(document, "P2 留存正例的独立场景图包含 45 个有效匹配、41 个内点。检测框坐标为 [90.4, 159.53, 285.45, 296.16]，浏览器框与报告 bbox 已逐项比对。图像来自 OpenCV 公开样例，不能称为团队家庭搬移实验。")
    text(document, "来源：docs/evidence/issue-3/browser/moved-matched.png 与对应 moved-report.json。登记使用 box.png 自匹配，仅作登记检查；box_in_scene.png 才是独立跨图正例。图片原始拍摄时间未知，A/B/C 与时间来自 Demo 编排。")

    document.add_page_break()
    heading(document, "5 失败与重放保护")
    text(document, "未确认负例检查 A、B、C 后得到 SUSPECTED，确认地点为空。旧图重放场景拒绝已见 B/A 图像，最终 UNKNOWN，只在新的 C 图像执行当前观察。两者都不能被解说成已找到或已证明不存在。")
    picture(document, "replay-rejected-B.png", "图 2  重放拒绝事件与历史证据分离  公开实拍照片  模拟导航", width=14.2)
    table(document, ["场景", "记忆优先步数", "固定顺序步数", "最终状态"], [
        ["正例 moved", "2", "2", "FOUND / B"],
        ["负例 missing", "3", "3", "SUSPECTED / null"],
        ["重放 replay", "3", "3", "UNKNOWN / null"],
    ], [4.0, 4.0, 4.0, 4.6])
    text(document, "对照冻结同一输入图片并分别登记运行。尝试步数打平，现有小样本不能支持搜索效率提升的结论；可证明的是执行链可见、结果可追溯。数据来源：docs/evidence/issue-3/comparison.json。")

    document.add_page_break()
    heading(document, "6 团队家庭图片适用性")
    text(document, "本地数据覆盖纸巾主案例及耳机壳、钥匙、笔、保温杯等备用物品。原图留在 local_data，不进入公开 GitHub 或本说明书。纸巾具有独立参考图、历史 A、当前 A、当前 B 和负例。其他物品没有完整历史序列时，只评估图片匹配，不虚构搬移过程。")
    table(document, ["对象与算法", "有效匹配 / 内点", "几何支持", "结论"], [
        ["纸巾历史 A  ORB", "21 / 7", "比例 0.3333\n覆盖 0.0524", "未确认"],
        ["纸巾当前 B  ORB", "15 / 3", "比例 0.2000\n覆盖 0.0012", "未确认"],
        ["纸巾当前 B  SIFT", "74 / 23", "比例 0.3108\n覆盖 0.0391", "未确认"],
        ["保温杯  AKAZE", "36 / 33", "比例 0.9167\n覆盖 0.0342", "未确认"],
    ], [5.0, 3.6, 4.4, 3.6])
    text(document, "上述匹配数量不是识别成功数量。纸巾新参考图清晰度指标为 102.53，局部匹配仍集中在不足的区域。保温杯虽然内点比例较高，但覆盖 0.0342 低于 0.10 的原判据，仍然不能确认。耳机壳参考图清晰度 21.265，质量门槛拒绝；钥匙和笔未确认，玻璃杯没有正例。")
    text(document, "有限对比共 25 个图对，分别执行 ORB、SIFT、AKAZE，形成 75 次检测评估：24 次带目标评价、51 次无目标评价，均无 FOUND。这是相同图片跨算法复评，不是 75 张独立照片，也不构成泛化准确率。生产检测器保留 ORB，不为当前照片放宽阈值。")
    text(document, "纸巾完整 SQLite 寻物中，moved 和 missing 均为 SUSPECTED 且确认地点为空，重放为 UNKNOWN。家庭页面展示这些实际未确认结果，不能称为家庭目标成功寻回。后续可改善参考图和拍摄条件，再以相同负例与原阈值复测。")
    text(document, "指标与逐图失败原因以 docs/TEAM_PHOTO_EVALUATION.md 和 docs/evidence/p3-photos-summary.json 为准；检测耗时属于本机离线照片匹配，不能代表端侧或机器人运动性能。")

    document.add_page_break()
    heading(document, "7 运行与测试")
    text(document, "Windows 环境使用 Python 3.11 或更高版本。以下命令在源码根目录执行，OpenCV 无需额外下载模型权重。新一次验证采用新的输出目录，避免覆盖既有证据。")
    for command in ["python -m venv .venv", '.\\.venv\\Scripts\\python.exe -m pip install -e ".[dev,vision]"', ".\\.venv\\Scripts\\python.exe -m pytest -q", ".\\.venv\\Scripts\\python.exe -m hometrace demo --port 8765", "浏览器打开 http://127.0.0.1:8765"]:
        p = text(document, command)
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(9)
    text(document, "界面依次运行寻找目标、未确认负例、旧图重放保护；单步检查当前照片、模拟导航与最终状态，下载 report.json 核对 location 与证据摘要。家庭案例使用本地私有配置入口，检测失败应继续显示未确认状态。")
    table(document, ["验证层级", "证据与判定"], [
        ["P0/P1/P2 历史基线", "P2 留存 148 passed，0 skipped；包含 P0 54、P1 68 和 P2 26 项"],
        ["P3 本轮回归", f"{test_count} passed；{test_log}" if test_count is not None else "本轮完整数量以 RUN_TEST_REPORT.md 与 PR 实际日志为准"],
        ["图像与证据", "逐图匹配指标、负例、重复图片保护、bbox 坐标和来源记录"],
        ["打包", "缺失/空文件/不可读材料拒绝；重开 ZIP 并核对每项 SHA256"],
    ], [4.4, 12.2])
    text(document, "148 项是已审核 P2 基线，不可替代最新代码的回归结果。Windows 全新环境安装、P3 最新测试及打包的实际运行日志见 docs/submission/RUN_TEST_REPORT.md。macOS/Linux、端侧设备与机器人未作本轮实机验收。")
    heading(document, "8 参赛交付核对")
    text(document, "团队自编技术说明书、技术展示 PPT、实际运行 MP4、核心源码、运行测试报告和开源说明共同构成材料集。个人身份与联系方式只在本地填写；打包工具从本地输入读取正式材料和团队信息，排除私图、数据库、密钥、缓存与虚拟环境。")
    text(document, "可读取文件与可提交内容是不同门槛。最终版本还需 PM 审查，团队核对身份与赛事平台要求，重新打开 ZIP、验证 SHA256 并人工解压查看。代码提交、打包成功或本说明书均不代表赛事上传已完成。")

    document.add_page_break()
    heading(document, "9 开源技术与原创性")
    text(document, "团队贡献包括 SQLite 时空记忆与证据模型、历史优先的搜索编排、当前证据与历史分离、旧图重放保护、明确不确定状态、逐步展示和可追溯结果报告。OpenCV 的 ORB、Hamming 匹配和 RANSAC 是复用的成熟开源方法，不作为从零原创的视觉算法。")
    table(document, ["组件", "使用方式与许可"], [
        ["Python / SQLite", "标准库 CLI、HTTP、JSON、ZIP 与数据库；PSF License / SQLite 公共领域"],
        ["OpenCV / NumPy", "特征与几何验证、数值数组；Apache-2.0 / BSD-3-Clause，保留随包告知"],
        ["pytest", "自动化回归，MIT，仅开发依赖"],
        ["公开图片", "OpenCV 4.12.0 samples/data，保留来源、SHA256 和许可文件"],
    ], [4.2, 12.4])
    text(document, "图片上游版本为 49486f61fb25722cbcf586b7f4320921d46fb38e，原图字节保持不变。来源与完整告知见 examples/recorded_images/sources.json、LICENSE.opencv.txt、COPYRIGHT.opencv.txt 和 LICENSE_CHANGE_NOTICE.opencv.txt。实际截图来自该输入驱动的软件运行，导航为模拟。")
    heading(document, "10 当前能力边界与后续验证")
    text(document, "当前适合演示有纹理平面参考外观的离线匹配、记忆查询与搜索验证逻辑。家庭小物体的尺度、视角、遮挡、反光、纹理不足和背景占比会使当前视觉方法失效。更换成熟特征方法不保证通过几何验证，必须同时核查负例。")
    text(document, "作品没有现场摄像头连续识别、实时建图、自主避障或真机控制，没有端侧运行速度、能耗或家庭泛化准确率。搜索效率尚未证明优于固定顺序。下一步以更合适的独立采集、相同验证门槛和明确设备实验补齐证据，不扩大非必要算法或框架。")
    heading(document, "证据索引")
    for item in ["GitHub 仓库  https://github.com/Qcx526/HomeTrace", "P2 稳定证据  docs/evidence/issue-3/", "家庭适用性  docs/TEAM_PHOTO_EVALUATION.md", "本轮运行测试  docs/submission/RUN_TEST_REPORT.md", "开源与原创性  docs/submission/OPEN_SOURCE_AND_ORIGINALITY.md", "提交材料清单  docs/submission/MATERIALS_CHECKLIST.md"]:
        text(document, item)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created team_authored DOCX: {output}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "docs/submission/HomeTrace_作品设计与技术说明书.docx")
    parser.add_argument("--test-count", type=int)
    parser.add_argument("--test-log")
    args = parser.parse_args()
    build(args.output.resolve(), args.test_count, args.test_log)


if __name__ == "__main__":
    main()
